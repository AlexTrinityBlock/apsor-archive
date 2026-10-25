"""File Library orchestration: multimodal intake, chunk reads, and paths."""

import asyncio
import hashlib
import json
import logging
import mimetypes
import zipfile
from collections.abc import Awaitable, Callable, Sequence
from dataclasses import dataclass
from datetime import datetime
from io import BytesIO
from typing import Any, TypeVar, cast
from uuid import uuid4

import pypdfium2 as pdfium
from docx import Document
from pypdf import PdfReader

from app.module.knowledge.domain.chunking_domain import chunk_text
from app.module.knowledge.domain.path_domain import (
    InvalidPathValueError,
    extension,
    file_name,
    normalize_directory_path,
    normalize_file_path,
    normalize_fuzzy_search,
    parent_directory,
    path_segments,
)
from app.module.knowledge.exception.domain_error import (
    AiProviderError,
    EmptyContentError,
    FileContentUnavailableError,
    FileRecordNotFoundError,
    FileTooLargeError,
    FileVersionConflictError,
    InvalidChunkWindowError,
    InvalidPathError,
    PathNotFoundError,
    TooManyChunksError,
    TreePathConflictError,
    UnsupportedFileTypeError,
    VisionExtractionError,
)
from app.module.knowledge.model.chunk_model import SOURCE_FILE, ChunkModel
from app.module.knowledge.model.file_model import FileModel
from app.module.knowledge.model.folder_model import FolderModel
from app.module.knowledge.protocol.cache_protocol import CacheProtocol
from app.module.knowledge.protocol.chunk_repository_protocol import ChunkRepositoryProtocol
from app.module.knowledge.protocol.file_repository_protocol import FileRepositoryProtocol
from app.module.knowledge.protocol.file_storage_protocol import FileStorageProtocol
from app.module.knowledge.protocol.folder_repository_protocol import FolderRepositoryProtocol
from app.module.knowledge.protocol.ingest_progress_protocol import IngestProgressProtocol
from app.module.knowledge.protocol.text_generator_protocol import TextGeneratorProtocol
from app.module.knowledge.protocol.vision_extractor_protocol import (
    DocumentAnalysis,
    VisionExtractorProtocol,
)
from app.module.knowledge.schema.file_schema import (
    ChunkSchema,
    DirectorySchema,
    FileChunksSchema,
    FileSchema,
    LibraryDeleteResultSchema,
    PathFileSchema,
    PathListingSchema,
)
from app.module.knowledge.service.indexer_service import IndexerService
from app.shared.time import utcnow

logger = logging.getLogger(__name__)

IngestCheckpoint = Callable[[], Awaitable[None]]
CpuRunner = Callable[..., Awaitable[Any]]
T = TypeVar("T")

IMAGE_EXTENSIONS = {".jpg", ".jpeg", ".png", ".webp"}
DOCUMENT_EXTENSIONS = {".pdf", ".docx"}
MARKDOWN_EXTENSIONS = {".md", ".markdown"}
TEXT_EXTENSIONS = {".md", ".markdown", ".txt", ".csv", ".json", ".log"}
ALLOWED_EXTENSIONS = TEXT_EXTENSIONS | DOCUMENT_EXTENSIONS | IMAGE_EXTENSIONS

DOCUMENT_ANALYSIS_PROMPT = (
    "You analyze a document before it enters a personal knowledge library. "
    "Return JSON with the keys: has_meaningful_text (boolean), use_vision "
    "(boolean), document_type (short string such as pdf, docx, text, json, "
    "image, scanned_pdf), title (short string), summary (1 to 2 sentences), tags "
    "(3 to 6 short strings), and language. Set use_vision to true only when the "
    "extracted text is empty, garbled, or clearly not the document's main content "
    "and a vision pass over its pages or images could recover it. For plain text "
    "files use_vision must be false. Do not invent content."
)


@dataclass(frozen=True)
class ExtractedDocument:
    """Text plus structured metadata produced by one extraction pass."""

    text: str
    document_type: str
    extraction_method: str
    ai_summary: str | None
    ai_tags: list[str]


@dataclass(frozen=True)
class FileOriginal:
    """The stored bytes and validators for one file version."""

    data: bytes
    file_name: str
    mime_type: str
    content_hash: str
    updated_at: datetime


class FileService:
    """Coordinate file storage, multimodal extraction, folder placement, and indexing."""

    def __init__(
        self,
        *,
        file_repository: FileRepositoryProtocol,
        folder_repository: FolderRepositoryProtocol,
        chunk_repository: ChunkRepositoryProtocol,
        indexer: IndexerService,
        storage: FileStorageProtocol,
        cache: CacheProtocol,
        text_generator: TextGeneratorProtocol,
        vision_extractor: VisionExtractorProtocol,
        target_chars: int,
        overlap_chars: int,
        max_chunks: int,
        max_upload_bytes: int,
        max_text_bytes: int,
        vision_max_pages: int,
        pdf_text_threshold: int,
        cpu_runner: CpuRunner | None = None,
    ) -> None:
        self.file_repository = file_repository
        self.folder_repository = folder_repository
        self.chunk_repository = chunk_repository
        self.indexer = indexer
        self.storage = storage
        self.cache = cache
        self.text_generator = text_generator
        self.vision_extractor = vision_extractor
        self.target_chars = target_chars
        self.overlap_chars = overlap_chars
        self.max_chunks = max_chunks
        self.max_upload_bytes = max_upload_bytes
        self.max_text_bytes = max_text_bytes
        self.vision_max_pages = vision_max_pages
        self.pdf_text_threshold = pdf_text_threshold
        self.cpu_runner: CpuRunner = cpu_runner or asyncio.to_thread

    async def _run_cpu(
        self,
        func: Callable[..., T],
        *args: Any,
        **kwargs: Any,
    ) -> T:
        """Run parsing or chunking on the caller-supplied bounded executor."""
        return cast(T, await self.cpu_runner(func, *args, **kwargs))

    @staticmethod
    def _to_schema(file: FileModel) -> FileSchema:
        return FileSchema(
            id=file.id,
            raw_file_path=file.raw_file_path,
            file_name=file.file_name,
            extension=file.extension,
            mime_type=file.mime_type,
            size_bytes=file.size_bytes,
            version=file.version,
            supersedes_file_id=file.supersedes_file_id,
            is_current=file.is_current,
            total_chunks=file.total_chunks,
            has_original=file.storage_key is not None,
            ai_summary=file.ai_summary,
            ai_tags=list(file.ai_tags),
            document_type=file.document_type,
            extraction_method=file.extraction_method,
            created_at=file.created_at,
            updated_at=file.updated_at,
            deleted_at=file.deleted_at,
        )

    @staticmethod
    def _path_file_schema(file: FileModel) -> PathFileSchema:
        return PathFileSchema(
            file_id=file.id,
            raw_file_path=file.raw_file_path,
            file_name=file.file_name,
            total_chunks=file.total_chunks,
            size_bytes=file.size_bytes,
            has_original=file.storage_key is not None,
            update_time=file.updated_at,
        )

    @staticmethod
    def _normalized_path(raw: str) -> str:
        try:
            return normalize_file_path(raw)
        except InvalidPathValueError as exc:
            raise InvalidPathError(str(exc)) from exc

    @classmethod
    def normalize_path(cls, raw: str) -> str:
        """Validate and normalize one logical file path."""
        return cls._normalized_path(raw)

    @staticmethod
    def _normalized_directory(raw: str) -> str:
        try:
            normalized = normalize_directory_path(raw)
        except InvalidPathValueError as exc:
            raise InvalidPathError(str(exc)) from exc
        if normalized == "/":
            raise InvalidPathError("A folder path must include a folder name.")
        return normalized

    @staticmethod
    def _directory_schema(folder: FolderModel) -> DirectorySchema:
        return DirectorySchema(path=folder.path, name=folder.name, depth=folder.depth)

    @staticmethod
    def _new_storage_key(suffix: str) -> str:
        return f"files/{uuid4().hex[:2]}/{uuid4().hex}{suffix}"

    @classmethod
    def new_storage_key(cls, suffix: str) -> str:
        """Return an opaque key for a new stored original."""
        return cls._new_storage_key(suffix)

    @staticmethod
    def _pdf_text(data: bytes) -> str:
        """Extract PDF text synchronously; callers run this in a thread."""
        try:
            reader = PdfReader(BytesIO(data))
            pages = [page.extract_text() or "" for page in reader.pages]
            return "\n\n".join(page for page in pages if page.strip())
        except Exception:
            return ""

    @staticmethod
    def _decode_text(data: bytes) -> str:
        if b"\x00" in data[:8192]:
            raise EmptyContentError()
        for encoding in ("utf-8", "utf-8-sig", "big5"):
            try:
                return data.decode(encoding)
            except UnicodeDecodeError:
                continue
        return data.decode("latin-1")

    @staticmethod
    def _heuristic_analysis(
        text: str,
        *,
        file_name_value: str,
        document_type: str,
    ) -> DocumentAnalysis:
        cleaned = text.strip()
        return DocumentAnalysis(
            has_meaningful_text=len(cleaned) >= 40,
            use_vision=False,
            document_type=document_type,
            title=file_name_value,
            summary=cleaned[:240],
            tags=[],
            language="",
        )

    async def _analyze_text(
        self,
        text: str,
        *,
        file_name_value: str,
        document_type: str,
    ) -> DocumentAnalysis:
        """Use structured AI output to judge whether the text is usable."""
        if not self.text_generator.configured or not text.strip():
            return self._heuristic_analysis(
                text,
                file_name_value=file_name_value,
                document_type=document_type,
            )
        user_prompt = (
            f"Filename: {file_name_value}\n"
            f"Declared type: {document_type}\n"
            f"Extracted text sample:\n{text[:5000]}"
        )
        try:
            completion = await self.text_generator.complete_json(
                DOCUMENT_ANALYSIS_PROMPT,
                user_prompt,
            )
            body = json.loads(completion.text)
        except (AiProviderError, KeyError, TypeError, ValueError) as exc:
            logger.warning("Document analysis failed; using heuristics: %s", exc)
            return self._heuristic_analysis(
                text,
                file_name_value=file_name_value,
                document_type=document_type,
            )
        raw_tags = body.get("tags")
        tags = (
            [str(tag).strip() for tag in raw_tags if str(tag).strip()]
            if isinstance(raw_tags, list)
            else []
        )
        return DocumentAnalysis(
            has_meaningful_text=bool(body.get("has_meaningful_text", len(text.strip()) >= 40)),
            use_vision=bool(body.get("use_vision", False)),
            document_type=str(body.get("document_type") or document_type),
            title=str(body.get("title") or file_name_value),
            summary=str(body.get("summary") or text[:240]),
            tags=tags[:6],
            language=str(body.get("language") or ""),
        )

    @staticmethod
    def _render_pdf_pages(data: bytes, max_pages: int) -> list[bytes]:
        """Render PDF pages to PNG bytes; callers run this in a thread."""
        images: list[bytes] = []
        pdf = pdfium.PdfDocument(data)
        try:
            for index in range(min(len(pdf), max_pages)):
                bitmap = pdf[index].render(scale=2.0)
                image = bitmap.to_pil()
                buffer = BytesIO()
                image.save(buffer, format="PNG")
                images.append(buffer.getvalue())
        finally:
            pdf.close()
        return images

    async def _vision_pdf(self, data: bytes) -> tuple[str, str | None, list[str], str]:
        """Render up to N PDF pages and OCR them with the vision provider."""
        texts: list[str] = []
        summary: str | None = None
        tags: list[str] = []
        document_type = "scanned_pdf"
        try:
            pages = await self._run_cpu(
                self._render_pdf_pages,
                data,
                self.vision_max_pages,
            )
            for page in pages:
                vision = await self.vision_extractor.extract(page, "image/png")
                if vision.text.strip():
                    texts.append(vision.text.strip())
                elif vision.caption.strip():
                    texts.append(vision.caption.strip())
                summary = summary or vision.summary or None
                tags.extend(vision.tags)
                document_type = vision.document_type or document_type
        except VisionExtractionError:
            raise
        except Exception as exc:
            raise EmptyContentError() from exc
        return "\n\n".join(texts).strip(), summary, list(dict.fromkeys(tags))[:8], document_type

    @staticmethod
    def _docx_text(data: bytes) -> str:
        document = Document(BytesIO(data))
        parts = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
        for table in document.tables:
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells]
                if any(cells):
                    parts.append(" | ".join(cells))
        return "\n".join(parts)

    @staticmethod
    def _docx_images(data: bytes) -> list[tuple[bytes, str]]:
        images: list[tuple[bytes, str]] = []
        with zipfile.ZipFile(BytesIO(data)) as archive:
            for name in archive.namelist():
                if not name.startswith("word/media/"):
                    continue
                suffix = extension(name)
                if suffix not in IMAGE_EXTENSIONS:
                    continue
                mime = {
                    ".jpg": "image/jpeg",
                    ".jpeg": "image/jpeg",
                    ".png": "image/png",
                    ".webp": "image/webp",
                }[suffix]
                images.append((archive.read(name), mime))
        return images

    async def _vision_images(
        self,
        images: list[tuple[bytes, str]],
    ) -> tuple[str, str | None, list[str], str]:
        texts: list[str] = []
        summary: str | None = None
        tags: list[str] = []
        document_type = "docx"
        for data, mime_type in images[: self.vision_max_pages]:
            vision = await self.vision_extractor.extract(data, mime_type)
            if vision.text.strip():
                texts.append(vision.text.strip())
            elif vision.caption.strip():
                texts.append(vision.caption.strip())
            summary = summary or vision.summary or None
            tags.extend(vision.tags)
            document_type = vision.document_type or document_type
        return "\n\n".join(texts).strip(), summary, list(dict.fromkeys(tags))[:8], document_type

    async def _extract_document(self, path: str, data: bytes) -> ExtractedDocument:
        """Extract text; use vision when text is absent or AI says it is needed."""
        suffix = extension(path)
        file_name_value = file_name(path)
        if suffix in TEXT_EXTENSIONS:
            text = await self._run_cpu(self._decode_text, data)
            analysis = await self._analyze_text(
                text,
                file_name_value=file_name_value,
                document_type=suffix.lstrip(".") or "text",
            )
            return ExtractedDocument(
                text=text,
                document_type=analysis.document_type,
                extraction_method="text",
                ai_summary=analysis.summary or None,
                ai_tags=analysis.tags,
            )
        if suffix == ".pdf":
            text = await self._run_cpu(self._pdf_text, data)
            analysis = await self._analyze_text(
                text,
                file_name_value=file_name_value,
                document_type="pdf",
            )
            needs_vision = (
                not analysis.has_meaningful_text
                or analysis.use_vision
                or len(text.strip()) < self.pdf_text_threshold
            )
            if needs_vision and self.vision_extractor.configured:
                vision_text, summary, tags, document_type = await self._vision_pdf(data)
                combined = "\n\n".join(part for part in (text.strip(), vision_text) if part)
                if not combined:
                    raise EmptyContentError()
                return ExtractedDocument(
                    text=combined,
                    document_type=document_type or "scanned_pdf",
                    extraction_method="mixed" if text.strip() else "vision",
                    ai_summary=summary or analysis.summary or None,
                    ai_tags=list(dict.fromkeys([*analysis.tags, *tags]))[:8],
                )
            if not text.strip():
                raise EmptyContentError()
            return ExtractedDocument(
                text=text,
                document_type=analysis.document_type,
                extraction_method="text",
                ai_summary=analysis.summary or None,
                ai_tags=analysis.tags,
            )
        if suffix == ".docx":
            try:
                text = await self._run_cpu(self._docx_text, data)
            except Exception as exc:
                raise EmptyContentError() from exc
            analysis = await self._analyze_text(
                text,
                file_name_value=file_name_value,
                document_type="docx",
            )
            if (
                not analysis.has_meaningful_text or analysis.use_vision
            ) and self.vision_extractor.configured:
                images = await self._run_cpu(self._docx_images, data)
                if images:
                    vision_text, summary, tags, document_type = await self._vision_images(images)
                    combined = "\n\n".join(part for part in (text.strip(), vision_text) if part)
                    if combined:
                        return ExtractedDocument(
                            text=combined,
                            document_type=document_type or "docx",
                            extraction_method="mixed" if text.strip() else "vision",
                            ai_summary=summary or analysis.summary or None,
                            ai_tags=list(dict.fromkeys([*analysis.tags, *tags]))[:8],
                        )
            if not text.strip():
                raise EmptyContentError()
            return ExtractedDocument(
                text=text,
                document_type=analysis.document_type,
                extraction_method="text",
                ai_summary=analysis.summary or None,
                ai_tags=analysis.tags,
            )
        if suffix in IMAGE_EXTENSIONS:
            if not self.vision_extractor.configured:
                # Local-first uploads must not depend on a vision provider. The
                # filename still gives the image a useful searchable caption.
                stem = file_name_value.rsplit(".", 1)[0]
                label = " ".join(stem.replace("_", " ").replace("-", " ").split())
                return ExtractedDocument(
                    text=f"{label or 'Untitled'} image",
                    document_type="image",
                    extraction_method="filename",
                    ai_summary=None,
                    ai_tags=[],
                )
            vision = await self.vision_extractor.extract(
                data,
                {
                    ".jpg": "image/jpeg",
                    ".jpeg": "image/jpeg",
                    ".png": "image/png",
                    ".webp": "image/webp",
                }[suffix],
            )
            text = vision.text.strip() or vision.caption.strip()
            if not text:
                raise EmptyContentError()
            return ExtractedDocument(
                text=text,
                document_type=vision.document_type or "image",
                extraction_method="vision",
                ai_summary=vision.summary or None,
                ai_tags=vision.tags,
            )
        raise UnsupportedFileTypeError(suffix)

    async def _ensure_folder_chain(self, directory: str) -> FolderModel | None:
        """Create every missing folder in ``directory`` and return the leaf folder."""
        parent_id: int | None = None
        current: FolderModel | None = None
        segments = path_segments(directory)
        for index, segment in enumerate(segments, start=1):
            directory_path = "/" + "/".join(segments[:index])
            folder = await self.folder_repository.get_by_path(directory_path)
            if folder is None:
                if await self.file_repository.get_current_by_path(directory_path) is not None:
                    raise TreePathConflictError(directory_path)
                folder = await self.folder_repository.add_if_absent(
                    FolderModel(
                        parent_id=parent_id,
                        path=directory_path,
                        name=segment,
                        depth=index,
                    )
                )
            parent_id = folder.id
            current = folder
        return current

    async def _ensure_parent_folder(self, path: str) -> int | None:
        folder = await self._ensure_folder_chain(parent_directory(path))
        return folder.id if folder is not None else None

    async def _ingest(
        self,
        *,
        path: str,
        data: bytes,
        text: str,
        mime_type: str,
        storage_key: str | None,
        occurred_at: datetime | None,
        ai_summary: str | None = None,
        ai_tags: list[str] | None = None,
        document_type: str = "text",
        extraction_method: str = "text",
        publish: bool = True,
        progress: IngestProgressProtocol | None = None,
        checkpoint: IngestCheckpoint | None = None,
    ) -> FileSchema:
        """Chunk, stage or publish a file version, and index its chunks.

        The synchronous API keeps ``publish=True``: old rows stop being current
        before the new version is inserted. The durable worker passes
        ``publish=False`` so a crash or embedding failure never exposes a
        partial file; ``publish_pending`` flips the staged rows visible once
        the job succeeds.
        """
        drafts = await self._run_cpu(
            chunk_text,
            text,
            target_chars=self.target_chars,
            overlap_chars=self.overlap_chars,
        )
        if not drafts:
            raise EmptyContentError()
        if len(drafts) > self.max_chunks:
            raise TooManyChunksError(len(drafts), self.max_chunks)
        if progress is not None:
            await progress.chunked(len(drafts))

        existing = await self.file_repository.get_current_by_path(path)
        if await self.folder_repository.get_by_path(path) is not None:
            raise TreePathConflictError(path)

        if publish and existing is not None:
            await self.file_repository.mark_not_current(existing)
            await self.chunk_repository.mark_source_not_current(SOURCE_FILE, existing.id)

        timestamp = occurred_at or utcnow()
        file = await self.file_repository.add(
            FileModel(
                raw_file_path=path,
                file_name=file_name(path),
                extension=extension(path),
                mime_type=mime_type,
                size_bytes=len(data),
                content_hash=hashlib.sha256(data).hexdigest(),
                storage_key=storage_key,
                ai_summary=ai_summary,
                ai_tags=ai_tags or [],
                document_type=document_type,
                extraction_method=extraction_method,
                version=existing.version + 1 if existing is not None else 1,
                supersedes_file_id=existing.id if existing is not None else None,
                is_current=publish,
                total_chunks=len(drafts),
                created_at=timestamp,
                updated_at=timestamp,
            )
        )

        await self._ensure_parent_folder(path)

        if checkpoint is not None:
            await checkpoint()
        if progress is not None:
            await progress.file_created(file.id)
        await self.indexer.index_chunks(
            source_kind=SOURCE_FILE,
            source_id=file.id,
            drafts=drafts,
            is_current=publish,
            progress=progress.embedded if progress is not None else None,
            checkpoint=checkpoint,
        )
        if publish:
            await self.cache.bump_version()
        return self._to_schema(file)

    async def ingest_extracted(
        self,
        *,
        path: str,
        data: bytes,
        extracted: ExtractedDocument,
        mime_type: str,
        storage_key: str | None,
        occurred_at: datetime | None,
        publish: bool = True,
        progress: IngestProgressProtocol | None = None,
        checkpoint: IngestCheckpoint | None = None,
    ) -> FileSchema:
        """Persist an already-extracted document through the shared pipeline."""
        return await self._ingest(
            path=path,
            data=data,
            text=extracted.text,
            mime_type=mime_type,
            storage_key=storage_key,
            occurred_at=occurred_at,
            ai_summary=extracted.ai_summary,
            ai_tags=extracted.ai_tags,
            document_type=extracted.document_type,
            extraction_method=extracted.extraction_method,
            publish=publish,
            progress=progress,
            checkpoint=checkpoint,
        )

    async def publish_pending(self, file_id: int) -> FileSchema:
        """Make a staged file version and its chunks current at once."""
        file = await self.file_repository.get_by_id(file_id)
        if file is None:
            raise FileRecordNotFoundError(file_id)
        if file.is_current:
            return self._to_schema(file)

        existing = await self.file_repository.get_current_by_path(file.raw_file_path)
        if existing is not None and existing.id != file.id:
            await self.file_repository.mark_not_current(existing)
            await self.chunk_repository.mark_source_not_current(SOURCE_FILE, existing.id)

        await self._ensure_parent_folder(file.raw_file_path)
        now = utcnow()
        file.is_current = True
        file.updated_at = now
        await self.chunk_repository.mark_source_current(SOURCE_FILE, file.id)
        await self.file_repository.flush()
        await self.cache.bump_version()
        return self._to_schema(file)

    async def discard_pending(self, file_id: int) -> bool:
        """Delete an unpublished file and its chunks after a failed job.

        Returns true only when a staged row was removed. A published file is
        left untouched so a late worker error cannot erase a visible version.
        """
        file = await self.file_repository.get_by_id(file_id)
        if file is None or file.is_current:
            return False
        await self.chunk_repository.delete_source(SOURCE_FILE, file.id)
        await self.file_repository.delete(file)
        return True

    async def extract_upload(self, *, path: str, data: bytes) -> ExtractedDocument:
        """Validate and extract one uploaded document for a queued job."""
        normalized = self._normalized_path(path)
        suffix = extension(normalized)
        if suffix not in ALLOWED_EXTENSIONS:
            raise UnsupportedFileTypeError(suffix)
        if len(data) > self.max_upload_bytes:
            raise FileTooLargeError(len(data), self.max_upload_bytes)
        return await self._extract_document(normalized, data)

    async def prepare_job_document(
        self,
        *,
        source_kind: str,
        path: str,
        data: bytes,
    ) -> ExtractedDocument:
        """Extract the payload referenced by a durable ingestion job."""
        if source_kind == "text":
            text = await self._run_cpu(self._decode_text, data)
            return ExtractedDocument(
                text=text,
                document_type=extension(path).lstrip(".") or "text",
                extraction_method="text",
                ai_summary=None,
                ai_tags=[],
            )
        return await self.extract_upload(path=path, data=data)

    async def ingest_text(
        self,
        *,
        path: str,
        content: str,
        occurred_at: datetime | None = None,
    ) -> FileSchema:
        """Create or version a text file and index it."""
        normalized = self._normalized_path(path)
        encoded = content.encode("utf-8")
        if len(encoded) > self.max_text_bytes:
            raise FileTooLargeError(len(encoded), self.max_text_bytes)
        mime_type = mimetypes.guess_type(normalized)[0] or "text/plain"
        return await self._ingest(
            path=normalized,
            data=encoded,
            text=content,
            mime_type=mime_type,
            storage_key=None,
            occurred_at=occurred_at,
            document_type=extension(normalized).lstrip(".") or "text",
            extraction_method="text",
        )

    async def update_markdown_content(
        self,
        file_id: int,
        *,
        content: str,
        expected_version: int,
    ) -> FileSchema:
        """以新版本更新既有 Markdown 檔案，保留舊版本與原始 bytes。"""
        file = await self._require_current_file_for_update(file_id)
        if file.extension not in MARKDOWN_EXTENSIONS:
            raise UnsupportedFileTypeError(file.extension)
        if file.version != expected_version:
            raise FileVersionConflictError(file_id, expected_version, file.version)

        encoded_content = content.encode("utf-8")
        if len(encoded_content) > self.max_text_bytes:
            raise FileTooLargeError(len(encoded_content), self.max_text_bytes)

        storage_key = self._new_storage_key(file.extension)
        await self.storage.save(storage_key, encoded_content)
        mime_type = file.mime_type if file.mime_type.startswith("text/") else "text/markdown"
        try:
            return await self._ingest(
                path=file.raw_file_path,
                data=encoded_content,
                text=content,
                mime_type=mime_type,
                storage_key=storage_key,
                occurred_at=None,
                ai_summary=file.ai_summary,
                ai_tags=list(file.ai_tags),
                document_type=file.document_type,
                extraction_method="text",
            )
        except Exception:
            await self.storage.delete(storage_key)
            raise

    async def ingest_upload(
        self,
        *,
        path: str,
        data: bytes,
        content_type: str | None,
        occurred_at: datetime | None = None,
    ) -> FileSchema:
        """Create or version an uploaded file and index its extracted text."""
        normalized = self._normalized_path(path)
        suffix = extension(normalized)
        if suffix not in ALLOWED_EXTENSIONS:
            raise UnsupportedFileTypeError(suffix)
        if len(data) > self.max_upload_bytes:
            raise FileTooLargeError(len(data), self.max_upload_bytes)

        extracted = await self._extract_document(normalized, data)
        mime_type = (
            content_type or mimetypes.guess_type(normalized)[0] or "application/octet-stream"
        )
        storage_key = self._new_storage_key(suffix)
        await self.storage.save(storage_key, data)
        try:
            return await self._ingest(
                path=normalized,
                data=data,
                text=extracted.text,
                mime_type=mime_type,
                storage_key=storage_key,
                occurred_at=occurred_at,
                ai_summary=extracted.ai_summary,
                ai_tags=extracted.ai_tags,
                document_type=extracted.document_type,
                extraction_method=extracted.extraction_method,
            )
        except Exception:
            await self.storage.delete(storage_key)
            raise

    async def create_folder(self, path: str) -> DirectorySchema:
        """Create one folder plus missing ancestors, idempotently."""
        normalized = self._normalized_directory(path)
        occupied_file = await self.file_repository.get_current_by_path(normalized)
        if occupied_file is not None:
            raise TreePathConflictError(normalized)
        existing = await self.folder_repository.get_by_path(normalized)
        if existing is not None:
            return self._directory_schema(existing)

        folder = await self._ensure_folder_chain(normalized)
        if folder is None:  # pragma: no cover - the root path is rejected above
            raise InvalidPathError("A folder path must include a folder name.")
        await self.cache.bump_version()
        return self._directory_schema(folder)

    async def _require_current_file_for_update(self, file_id: int) -> FileModel:
        """以資料列鎖定 current file；供更新、移動、複製、刪除共用。"""
        file = await self.file_repository.get_current_for_update(file_id)
        if file is None or file.deleted_at is not None or not file.is_current:
            raise FileRecordNotFoundError(file_id)
        return file

    @staticmethod
    def _target_file_path(file: FileModel, raw: str) -> str:
        target = FileService._normalized_path(raw)
        if extension(target) != file.extension:
            raise InvalidPathError("A moved or copied file must keep its original extension.")
        return target

    async def _target_is_occupied(self, path: str) -> bool:
        if await self.folder_repository.get_by_path(path) is not None:
            return True
        return await self.file_repository.get_current_by_path(path) is not None

    def _merge_chunk_contents(self, chunks: list[ChunkModel]) -> str:
        """Rebuild text by removing the known overlap between adjacent chunks."""
        contents = [chunk.content for chunk in chunks if chunk.is_current]
        if not contents:
            return ""
        merged = contents[0]
        for content in contents[1:]:
            if not content:
                continue
            max_overlap = min(self.overlap_chars, len(merged), len(content))
            overlap = 0
            for size in range(max_overlap, 0, -1):
                if merged.endswith(content[:size]):
                    overlap = size
                    break
            if overlap:
                merged += content[overlap:]
                continue
            separator = (
                "" if merged.endswith(("\n", " ")) or content.startswith(("\n", " ")) else "\n"
            )
            merged = f"{merged}{separator}{content}"
        return merged

    async def move_file(self, file_id: int, *, path: str) -> FileSchema:
        """Move one current file in place, preserving its identity and chunks."""
        file = await self._require_current_file_for_update(file_id)
        target = self._target_file_path(file, path)
        if target == file.raw_file_path:
            return self._to_schema(file)
        if await self._target_is_occupied(target):
            raise TreePathConflictError(target)

        await self._ensure_parent_folder(target)
        await self.file_repository.update_location(
            file,
            path=target,
            file_name=file_name(target),
        )
        await self.cache.bump_version()
        return self._to_schema(file)

    async def rename_file(self, file_id: int, *, file_name_value: str) -> FileSchema:
        """Rename one current file without changing its parent folder or extension."""
        file = await self._require_current_file_for_update(file_id)
        name = file_name_value.strip()
        # 先擋掉 slash，否則 target 會被 normalize 成合法的多段路徑；其餘
        # 空名稱、"."、".." 與 control characters 由 normalize_file_path 擋下。
        if not name or "/" in name or "\\" in name:
            raise InvalidPathError("A renamed file must use a single non-empty file name.")

        parent = parent_directory(file.raw_file_path)
        target = f"/{name}" if parent == "/" else f"{parent}/{name}"
        target = self._target_file_path(file, target)
        if target == file.raw_file_path:
            return self._to_schema(file)
        if await self._target_is_occupied(target):
            raise TreePathConflictError(target)

        await self.file_repository.update_location(
            file,
            path=target,
            file_name=name,
        )
        await self.cache.bump_version()
        return self._to_schema(file)

    async def copy_file(self, file_id: int, *, path: str) -> FileSchema:
        """Copy one current file into a new file record with fresh chunks."""
        source = await self._require_current_file_for_update(file_id)
        target = self._target_file_path(source, path)
        if await self._target_is_occupied(target):
            raise TreePathConflictError(target)

        chunks = await self.chunk_repository.get_source_chunks(SOURCE_FILE, source.id)
        text = self._merge_chunk_contents(chunks)
        if not text.strip():
            raise EmptyContentError()

        storage_key: str | None = None
        if source.storage_key is not None:
            data = await self.storage.read(source.storage_key)
            storage_key = self._new_storage_key(source.extension)
            await self.storage.save(storage_key, data)
        else:
            data = text.encode("utf-8")

        try:
            return await self._ingest(
                path=target,
                data=data,
                text=text,
                mime_type=source.mime_type,
                storage_key=storage_key,
                occurred_at=None,
                ai_summary=source.ai_summary,
                ai_tags=list(source.ai_tags),
                document_type=source.document_type,
                extraction_method=source.extraction_method,
            )
        except Exception:
            if storage_key is not None:
                await self.storage.delete(storage_key)
            raise

    async def get_file(self, file_id: int) -> FileSchema:
        file = await self.file_repository.get_by_id(file_id)
        if file is None:
            raise FileRecordNotFoundError(file_id)
        return self._to_schema(file)

    async def read_chunks(
        self,
        file_id: int,
        *,
        chunk: int,
        chunk_range: int,
    ) -> FileChunksSchema:
        if chunk < 1:
            raise InvalidChunkWindowError("chunk must be at least 1.")
        if chunk_range < 0 or chunk_range > 25:
            raise InvalidChunkWindowError("chunk_range must be between 0 and 25.")
        file = await self.file_repository.get_by_id(file_id)
        if file is None:
            raise FileRecordNotFoundError(file_id)
        if file.total_chunks == 0:
            raise InvalidChunkWindowError("This file has no chunks.")
        if chunk > file.total_chunks:
            raise InvalidChunkWindowError(
                f"chunk {chunk} exceeds the file's {file.total_chunks} chunks."
            )
        start = max(1, chunk - chunk_range)
        end = min(file.total_chunks, chunk + chunk_range)
        chunks = await self.chunk_repository.get_chunk_window(
            SOURCE_FILE,
            file_id,
            start,
            end,
        )
        return FileChunksSchema(
            file_id=file.id,
            raw_file_path=file.raw_file_path,
            file_name=file.file_name,
            update_time=file.updated_at,
            total_chunks=file.total_chunks,
            chunk_start=start,
            chunk_end=end,
            chunks=[
                ChunkSchema(
                    chunk=item.chunk_index,
                    content=item.content,
                    char_count=item.char_count,
                    token_count=item.token_count,
                    is_current=item.is_current,
                )
                for item in chunks
            ],
        )

    async def list_path(
        self,
        *,
        path: str,
        path_depth: int,
        fuzzy_search: str | None,
    ) -> PathListingSchema:
        try:
            normalized = normalize_directory_path(path)
            fuzzy = normalize_fuzzy_search(fuzzy_search)
        except InvalidPathValueError as exc:
            raise InvalidPathError(str(exc)) from exc
        if path_depth < 1 or path_depth > 10:
            raise InvalidPathError("path_depth must be between 1 and 10.")

        if fuzzy is not None:
            files = await self.file_repository.search_paths(normalized, fuzzy, 50)
            return PathListingSchema(
                path=normalized,
                fuzzy_search=fuzzy,
                directories=[],
                files=[self._path_file_schema(file) for file in files],
            )

        folder = None if normalized == "/" else await self.folder_repository.get_by_path(normalized)
        if normalized != "/" and folder is None:
            raise PathNotFoundError(normalized)

        if path_depth == 1:
            children = await self.folder_repository.get_children(
                None if folder is None else folder.id
            )
        else:
            children = await self.folder_repository.list_descendants(
                root_path=normalized,
                root_depth=folder.depth if folder is not None else 0,
                max_depth=path_depth,
                limit=1000,
            )

        directories = [
            self._directory_schema(child) for child in children if child.path != normalized
        ]
        files = await self.file_repository.list_files_in_directory(normalized)
        return PathListingSchema(
            path=normalized,
            fuzzy_search=None,
            directories=directories,
            files=[self._path_file_schema(file) for file in files],
        )

    async def delete_file(self, file_id: int) -> None:
        # DELETE 只允許 current file；歷史版本 id 一律回 404，避免誤刪已取代的版本。
        file = await self._require_current_file_for_update(file_id)
        now = utcnow()
        file.is_current = False
        file.deleted_at = now
        file.updated_at = now
        await self.chunk_repository.mark_source_not_current(SOURCE_FILE, file.id)
        await self.cache.bump_version()

    async def delete_library_selection(
        self,
        *,
        file_ids: Sequence[int],
        folder_paths: Sequence[str],
    ) -> LibraryDeleteResultSchema:
        """原子地軟刪除選取的檔案，並遞迴軟刪除所選資料夾下的內容。

        資料夾只保留最上層路徑，子資料夾由父路徑遞迴涵蓋；folder rows 直接
        刪除，歷史 file rows 與原始 bytes 一律保留。
        """
        normalized_folders = list(
            dict.fromkeys(self._normalized_directory(path) for path in folder_paths)
        )
        unique_file_ids = list(dict.fromkeys(file_ids))

        # 先按深度排序，父資料夾一定排在子資料夾前面，再濾掉被涵蓋的路徑。
        top_level_folders: list[str] = []
        for path in sorted(normalized_folders, key=lambda value: len(path_segments(value))):
            if any(path.startswith(f"{parent}/") for parent in top_level_folders):
                continue
            top_level_folders.append(path)

        if not unique_file_ids and not top_level_folders:
            raise InvalidPathError("Select at least one file or folder.")

        files = await self.file_repository.list_current_for_update(
            file_ids=unique_file_ids,
            folder_paths=top_level_folders,
        )
        files_by_id = {file.id: file for file in files}
        for file_id in unique_file_ids:
            if file_id not in files_by_id:
                raise FileRecordNotFoundError(file_id)

        folders = await self.folder_repository.list_for_update(top_level_folders)
        found_paths = {folder.path for folder in folders}
        for path in top_level_folders:
            if path not in found_paths:
                raise PathNotFoundError(path)

        now = utcnow()
        for file in files:
            file.is_current = False
            file.deleted_at = now
            file.updated_at = now
        deleted_file_ids = sorted(files_by_id)
        await self.chunk_repository.mark_sources_not_current(SOURCE_FILE, deleted_file_ids)
        await self.folder_repository.delete_many(folders)
        await self.cache.bump_version()
        return LibraryDeleteResultSchema(
            deleted_file_ids=deleted_file_ids,
            deleted_folder_paths=top_level_folders,
            deleted_files=len(deleted_file_ids),
            deleted_folders=len(folders),
        )

    async def download(self, file_id: int) -> tuple[bytes, str, str]:
        original = await self.read_original(file_id)
        return original.data, original.file_name, original.mime_type

    async def read_original(self, file_id: int) -> FileOriginal:
        """Load the stored original for inline viewing or attachment download.

        Text created through ``POST /files/text`` has no stored original; its
        content is only available through the indexed chunks. That case keeps
        the existing ``file_content_unavailable`` domain error so callers can
        fall back explicitly instead of receiving an empty body.
        """
        file = await self.file_repository.get_by_id(file_id)
        if file is None:
            raise FileRecordNotFoundError(file_id)
        if not file.storage_key:
            raise FileContentUnavailableError(file_id)
        data = await self.storage.read(file.storage_key)
        return FileOriginal(
            data=data,
            file_name=file.file_name,
            mime_type=file.mime_type,
            content_hash=file.content_hash,
            updated_at=file.updated_at,
        )
